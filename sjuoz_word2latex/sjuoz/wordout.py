# -*- coding: utf-8 -*-
"""Hazırlanan makalenin dergi biçiminde Word (.docx) sürümü.

PDF'ten Word'e çevirmek (pdf2docx) iki sütunlu sayfada kelimeleri yapıştırıp düzeni bozduğu için
Word dosyası PDF'ten değil, aynı içerikten üretilir:

1. Yazıcının (Writer) ürettiği gövde LaTeX'i pandoc'un okuyabileceği sade LaTeX'e indirgenir
   (\\sjuozfit, minipage şekilleri, tabularx/xltabular tablolar, kaynakça listesi …).
2. pandoc bunu .docx'e çevirir: denklemler gerçek Word denklemi (OMML), tablolar Word tablosu olur.
3. python-docx ile dergi düzeni eklenir: ön kısım (makale türü, başlık, yazarlar, kurumlar, sorumlu
   yazar, tarihli özet kutusu, anahtar kelimeler), tek sütundan iki sütuna geçiş, ilk sayfa üst bilgisi
   (dergi adı, DOI, logo), diğer sayfalarda kısa yazar adı ve SJUOZ|VOL… satırı, altlıkta DOI ve
   sayfa numarası, Times New Roman yazı biçimleri.
"""
import os
import re

from .common import warn

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ---------------------------------------------------------------- LaTeX -> pandoc'a uygun LaTeX
def _group(s, i):
    """s[i] == '{' ise eşleşen '}' sonrası indeks ve içerik."""
    depth, j = 0, i
    while j < len(s):
        c = s[j]
        if c == "\\":
            j += 2
            continue
        if c == "{":
            depth += 1
        elif c == "}":
            depth -= 1
            if depth == 0:
                return j + 1, s[i + 1:j]
        j += 1
    return len(s), s[i + 1:]


def _replace_cmd(s, name, nargs, fn):
    """\\name{a}{b}… komutlarını fn(args) ile değiştirir (iç içe parantezlere dayanıklı)."""
    out, pos = [], 0
    pat = re.compile(r"\\%s(?![A-Za-z])\s*" % re.escape(name))
    while True:
        mt = pat.search(s, pos)
        if not mt:
            out.append(s[pos:])
            return "".join(out)
        j, args = mt.end(), []
        while len(args) < nargs and j < len(s) and s[j] == "{":
            j, arg = _group(s, j)
            args.append(arg)
            while j < len(s) and s[j] in " \t" and len(args) < nargs:
                j += 1
        if len(args) < nargs:
            out.append(s[pos:mt.end()])
            pos = mt.end()
            continue
        out.append(s[pos:mt.start()])
        out.append(fn(args))
        pos = j


COLS_MARK = {1: "SJUOZTEKSUTUN", 2: "SJUOZIKISUTUN"}   # pandoc sonrası bölüm sonuna çevrilir


def _figure(block, number, label, width="3.3in"):
    imgs = re.findall(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}", block)
    cap = ""
    mt = re.search(r"\\caption(?:of\{figure\})?\{", block)
    if mt:
        cap = _group(block, mt.end() - 1)[1]
    head = r"\textbf{%s %s:} " % (label, number) if label == "Figure" else r"\textbf{%s (%s):} " % (label, number)
    parts = [r"\begin{figure}"]
    for f in imgs:
        parts.append(r"\includegraphics[width=%s]{%s}" % (width, f))
    parts.append(r"\caption{%s}" % (head + cap) if cap else "")
    parts.append(r"\end{figure}")
    return "\n".join(p for p in parts if p)


def _rows(tab_body):
    """Tablo gövdesindeki satırlar (kural komutları atılır)."""
    body = re.sub(r"\\(toprule|midrule|bottomrule|hline)\b", "", tab_body)
    body = re.sub(r"\\addlinespace(\[[^\]]*\])?", "", body)
    body = re.sub(r"\\cmidrule(\([^)]*\))?\{[^}]*\}", "", body)
    rows = [r.strip() for r in re.split(r"\\\\(?:\[[^\]]*\])?", body)]
    return [r for r in rows if r]


def _table(caption, ncols, head_rows, body_rows, number, label, widths=None):
    lines = [r"\begin{table}"]
    if caption is not None:
        head = r"\textbf{%s %s:} " % (label, number) if label == "Table" else r"\textbf{%s (%s):} " % (label, number)
        lines.append(r"\caption{%s}" % (head + caption))
    widths = widths or [1.0] * ncols
    # Word'de yazı 9 pt: "99.84" gibi sayılar alt alta bölünmesin diye her sütun en az ~%8,5
    frac = [max(w / sum(widths), min(0.085, 1.0 / ncols)) for w in widths]
    total = sum(frac)
    spec = "".join(r"p{%.3f\linewidth}" % (0.98 * w / total) for w in frac)
    lines.append(r"\begin{tabular}{%s}" % spec)
    lines.append(r"\toprule")
    for r in head_rows:
        lines.append(r + r" \\")
    if head_rows:
        lines.append(r"\midrule")
    for r in body_rows:
        lines.append(r + r" \\")
    lines.append(r"\bottomrule")
    lines.append(r"\end{tabular}")
    lines.append(r"\end{table}")
    return "\n".join(lines)


def _colcount(spec):
    spec = re.sub(r"[@>]\{(?:[^{}]|\{[^{}]*\})*\}", "", spec)
    return max(1, len(re.findall(r"[lcrXp]", spec)))


def _blocks(s):
    """Şekil/tablo kutuları: "\\noindent\\begin{minipage}{\\linewidth} … \\end{minipage}" (iç içe minipage'lere
    dayanıklı) ve uzun tablolar için "\\begingroup … \\endgroup". [(başlangıç, bitiş, içerik)]"""
    out = []
    opener = re.compile(r"\\noindent\s*\\begin\{minipage\}\{\\linewidth\}|\\begingroup")
    pos = 0
    while True:
        mt = opener.search(s, pos)
        if not mt:
            return out
        if mt.group(0) == "\\begingroup":
            end = s.find("\\endgroup", mt.end())
            if end < 0:
                return out
            out.append((mt.start(), end + len("\\endgroup"), s[mt.end():end]))
            pos = end + len("\\endgroup")
            continue
        depth, j = 1, mt.end()
        tok = re.compile(r"\\(begin|end)\{minipage\}")
        while depth:
            t = tok.search(s, j)
            if not t:
                return out
            depth += 1 if t.group(1) == "begin" else -1
            j = t.end()
        end = j
        tail = re.match(r"\s*\\par\\vspace\{6pt\}", s[end:])
        out.append((mt.start(), end + (tail.end() if tail else 0), s[mt.end():j - len("\\end{minipage}")]))
        pos = end


def pandoc_latex(body):
    """Yazıcının gövde LaTeX'ini pandoc'un okuyabileceği sade LaTeX'e çevirir."""
    s = body
    fig_n = tab_n = eq_n = sec_n = 0
    fig_label, tab_label = "Figure", "Table"

    # denklem: numara metin olarak eklenir (Word'de numara sağda görünür)
    def eq(mt):
        nonlocal eq_n
        star, inner = mt.group(1), mt.group(2)
        tag = re.search(r"\\tag\{([^}]*)\}", inner)
        inner = re.sub(r"\\tag\{[^}]*\}", "", inner)
        num = None
        if not star:
            eq_n += 1
            num = tag.group(1) if tag else str(eq_n)
        inner = inner.strip()
        return "\n\\[ %s %s\\]\n" % (inner, (r"\qquad (%s)" % num) if num else "")
    s = _replace_cmd(s, "sjuozfitone", 1, lambda a: a[0])
    s = _replace_cmd(s, "sjuozfit", 3, lambda a: a[0])
    s = re.sub(r"\\begin\{equation(\*?)\}(.*?)\\end\{equation\*?\}", eq, s, flags=re.S)

    out, pos = [], 0
    for start, end, block in _blocks(s):
        out.append(s[pos:start])
        pos = end
        before = s[:start]
        in_cols = before.rfind("\\begin{multicols}") > before.rfind("\\end{multicols}")
        name = re.search(r"\\renewcommand\{\\(figure|table)name\}\{([^}]*)\}", block)
        if "\\includegraphics" in block and "tabular" not in block:
            fig_n += 1
            sub = re.search(r"\\renewcommand\{\\thefigure\}\{([^}]*)\}", block)
            num = sub.group(1) if sub and "\\arabic" not in sub.group(1) else str(fig_n)
            if sub and "\\arabic" not in sub.group(1):
                fig_n -= 0 if num.split("-")[-1] == "A" else 1
            imgs = len(re.findall(r"\\includegraphics", block))
            col_w = 3.3 if in_cols else 6.8
            out.append(_figure(block, num, name.group(2) if name else fig_label,
                               "%.2fin" % (col_w if imgs == 1 else max(1.5, col_w / min(imgs, 2) - 0.1))))
            continue
        tm = re.search(r"\\begin\{(tabularx|xltabular)\}\{[^}]*\}\{((?:[^{}]|\{(?:[^{}]|\{[^{}]*\})*\})*)\}(.*?)\\end\{\1\}",
                       block, re.S)
        if tm:
            tab_n += 1
            ncols = _colcount(tm.group(2))
            cap = None
            cm = re.search(r"\\caption(?:of\{table\})?\{", block)
            if cm:
                cap = _group(block, cm.end() - 1)[1]
                cap = re.sub(r"\\hypertarget\{[^}]*\}\{\}", "", cap)
            inner = tm.group(3)
            if tm.group(1) == "xltabular":
                first = inner.split(r"\endfirsthead")[0]
                first = re.sub(r"^.*?\\\\", "", first, count=1, flags=re.S)  # \caption satırı
                rest = inner.split(r"\endlastfoot")[-1]
                head_rows, body_rows = _rows(first), _rows(rest)
            else:
                rows = _rows(inner)
                n_head = 1
                hm = re.search(r"\\midrule", inner)
                if hm:
                    n_head = max(1, len(_rows(inner[:hm.start()])))
                head_rows, body_rows = rows[:n_head], rows[n_head:]
            hs = [float(x) for x in re.findall(r"\\hsize=([\d.]+)\\hsize", tm.group(2))]
            out.append(_table(cap, ncols, head_rows, body_rows, tab_n, name.group(2) if name else tab_label,
                              hs if len(hs) == ncols else None))
            continue
        out.append(block)
    out.append(s[pos:])
    s = "".join(out)

    # bölüm numaraları (Word'de otomatik numara yok)
    def sec(mt):
        nonlocal sec_n
        sec_n += 1
        return r"\section{%d. %s}" % (sec_n, mt.group(1))
    s = re.sub(r"\\section\{((?:[^{}]|\{[^{}]*\})*)\}", sec, s)

    # kaynakça listesi -> paragraflar
    def refs(mt):
        items = re.split(r"\\item\s+", mt.group(1))
        return "\n\n".join(i.strip() for i in items[1:] if i.strip())
    s = re.sub(r"\\begin\{list\}\{\}\{.*?\n\}\n(.*?)\\end\{list\}", refs, s, flags=re.S)

    s = _replace_cmd(s, "hyperlink", 2, lambda a: a[1])
    s = _replace_cmd(s, "hypertarget", 2, lambda a: a[1])
    s = re.sub(r"\{\\urlstyle\{same\}\\nolinkurl\{([^}]*)\}\}", r"\1", s)
    # sütun geçişleri: işaret paragrafı olarak bırakılır, Word'de bölüm sonuna çevrilir
    s = re.sub(r"\\begin\{multicols\}\{2\}", "\n\n%s\n\n" % COLS_MARK[1], s)
    s = re.sub(r"\\end\{multicols\}", "\n\n%s\n\n" % COLS_MARK[2], s)
    s = re.sub(r"\\(begin|end)\{landscape\}", "", s)
    s = re.sub(r"\\setlength\{\\premulticols\}\{[^}]*\}", "", s)
    s = re.sub(r"\\(setcounter|addtocounter)\{[^}]*\}\{[^}]*\}", "", s)
    s = _replace_cmd(s, "renewcommand", 2, lambda a: "")
    s = re.sub(r"\\captionsetup\{[^}]*\}", "", s)
    s = re.sub(r"\\(par)?\\?vspace\*?\{[^}]*\}", "", s)
    s = re.sub(r"\\(noindent|centering|clearpage|newpage|footnotesize|small|normalsize|par)\b", "", s)
    s = re.sub(r"\\begin\{rtlpar\}(.*?)\\end\{rtlpar\}", r"\1", s, flags=re.S)
    return s


# ---------------------------------------------------------------- Word düzeni
def _qn(tag):
    return "{%s}%s" % (W_NS, tag)


def _set_cols(sectPr, num, space=360):
    from lxml import etree
    cols = sectPr.find(_qn("cols"))
    if cols is None:
        cols = etree.SubElement(sectPr, _qn("cols"))
    cols.set(_qn("num"), str(num))
    cols.set(_qn("space"), str(space))


def _grid(table, widths):
    """Sütun genişlikleri (tblGrid + hücreler, sabit düzen): Word ve LibreOffice aynı göstersin."""
    from lxml import etree
    tbl = table._tbl
    grid = tbl.tblGrid
    for k, gc in enumerate(grid.findall(_qn("gridCol"))):
        gc.set(_qn("w"), str(int(widths[k].inches * 1440)))
    tblPr = tbl.tblPr
    lay = etree.SubElement(tblPr, _qn("tblLayout"))
    lay.set(_qn("type"), "fixed")
    tw = tblPr.find(_qn("tblW"))
    if tw is None:
        tw = etree.SubElement(tblPr, _qn("tblW"))
    tw.set(_qn("w"), str(int(sum(w.inches for w in widths) * 1440)))
    tw.set(_qn("type"), "dxa")
    for row in table.rows:
        for c, w in zip(row.cells, widths):
            c.width = w


def _shade(cell, hexcolor):
    from lxml import etree
    tcPr = cell._tc.get_or_add_tcPr()
    shd = etree.SubElement(tcPr, _qn("shd"))
    shd.set(_qn("val"), "clear")
    shd.set(_qn("color"), "auto")
    shd.set(_qn("fill"), hexcolor)


def _borders(table, top=None, bottom=None):
    """Tablonun yalnız üst/alt çizgisi (pt*8 birimi)."""
    from lxml import etree
    tblPr = table._tbl.tblPr
    b = etree.SubElement(tblPr, _qn("tblBorders"))
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = etree.SubElement(b, _qn(side))
        size = {"top": top, "bottom": bottom}.get(side)
        e.set(_qn("val"), "single" if size else "nil")
        if size:
            e.set(_qn("sz"), str(size))
            e.set(_qn("color"), "000000")


def _field(run, instr):
    from lxml import etree
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        if kind == "instr":
            el = etree.SubElement(run._r, _qn("instrText"))
            el.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            el.text = text
        elif kind == "text":
            el = etree.SubElement(run._r, _qn("t"))
            el.text = text
        else:
            el = etree.SubElement(run._r, _qn("fldChar"))
            el.set(_qn("fldCharType"), kind)


def _plain(latex):
    """Kısa LaTeX parçasından düz metin (başlık/yazar satırı için)."""
    s = _replace_cmd(latex, "textit", 1, lambda a: a[0])
    s = _replace_cmd(s, "textbf", 1, lambda a: a[0])
    s = s.replace(r"\&", "&").replace(r"\%", "%").replace(r"\_", "_").replace(r"\#", "#").replace("~", " ")
    return re.sub(r"\\[A-Za-z]+\*?", "", s).replace("{", "").replace("}", "").strip()


def _add_text(par, text, size=None, bold=False, italic=False, sup=False):
    r = par.add_run(text)
    r.bold, r.italic = bold or None, italic or None
    if size:
        from docx.shared import Pt
        r.font.size = Pt(size)
    if sup:
        r.font.superscript = True
    return r


def _italic_runs(par, tokens, upper, size, bold=True):
    """Başlık: italik tür adları korunur, gerisi büyük harf."""
    for t in tokens:
        if t[0] != "t" or not t[1]:
            continue
        txt = re.sub(r"\s+", " ", t[1])
        it = bool(t[2][1]) and bool(re.search(r"[a-z]", txt))
        _add_text(par, txt if it or not upper else txt.upper(), size, bold=bold, italic=it)


def build_front(doc, art, opts, meta_info):
    """Ön kısım paragrafları (belgenin başına eklenecek öğe listesi döndürür)."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, Inches
    from .template import keywords_case
    els = []

    def para(align=None, space_after=2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.first_line_indent = Pt(0)
        if align == "c":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif align == "j":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        els.append(p._p)
        return p

    _add_text(para(space_after=6), opts.article_type or art.article_type, 10, bold=True)
    p = para("c", 6)
    toks = getattr(art, "title_tokens", []) or [("t", art.title, (True, False, None, None))]
    _italic_runs(p, toks, not opts.keep_title_case, 12)
    p = para("c", 4)
    n = len(art.authors)
    for k, au in enumerate(art.authors):
        if k:
            _add_text(p, ", and " if k == n - 1 and n > 2 else (" and " if k == n - 1 else ", "), 10, bold=True)
        _add_text(p, au["name"] + " ", 10, bold=True)
        if au["marks"]:
            _add_text(p, ",".join(au["marks"]), 10, bold=True, sup=True)
        if au.get("orcid"):
            orcid_png = meta_info.get("orcid_png")
            if orcid_png and os.path.isfile(orcid_png):
                r = p.add_run()
                r.add_picture(orcid_png, width=Pt(8))
            _add_link(p, au["orcid"], " iD" if not (orcid_png and os.path.isfile(orcid_png)) else "")
    for k, a in enumerate(art.affiliations):
        p = para(space_after=0)
        _add_text(p, str(k + 1), 10, sup=True)
        _add_text(p, a, 10)
    p = para("c", 6)
    p.paragraph_format.space_before = Pt(6)
    _add_text(p, "*", 9, sup=True)
    _add_text(p, "Corresponding author, E-mail: ", 9, italic=True)
    if art.corr_email:
        _add_link(p, "mailto:" + art.corr_email, art.corr_email, color="0000FF")
    if art.corr_tel:
        _add_text(p, " (%s: %s)" % (getattr(art, "corr_tel_label", "Tel"), art.corr_tel), 9)

    # özet kutusu: solda tarihler, sağda gölgeli özet
    tbl = doc.add_table(rows=2, cols=2)
    tbl.autofit = False
    _borders(tbl, top=12, bottom=12)
    _grid(tbl, (Inches(1.25), Inches(5.75)))
    h = tbl.cell(0, 1).paragraphs[0]
    _add_text(h, "ABSTRACT", 10, bold=True)
    dates = tbl.cell(1, 0)
    dp = dates.paragraphs[0]
    for k, (lab, val) in enumerate((("Received", opts.received), ("Accepted", opts.accepted),
                                    ("Published", opts.published))):
        if k:
            dp = dates.add_paragraph()
        dp.paragraph_format.space_before = Pt(10 if k else 14)
        _add_text(dp, lab, 9)
        dp.add_run().add_break()
        _add_text(dp, val or "xxx", 9)
    ab = tbl.cell(1, 1)
    _shade(ab, "E2EAF4")
    first = True
    for txt in art.abstract_plain or [""]:
        ap = ab.paragraphs[0] if first else ab.add_paragraph()
        first = False
        ap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        ap.paragraph_format.first_line_indent = Pt(0)
        _add_text(ap, re.sub(r"\s+", " ", txt).strip(), 10)
    kp = ab.add_paragraph()
    kp.paragraph_format.space_before = Pt(3)
    _add_text(kp, "Keywords: ", 9, bold=True)
    kw = art.keywords.strip()
    kw_tokens = getattr(art, "keyword_tokens", [])
    if kw_tokens and any(t[2][1] for t in kw_tokens if t[0] == "t"):
        for t in kw_tokens:
            if t[0] == "t" and t[1]:
                _add_text(kp, t[1], 9, italic=bool(t[2][1]))
    else:
        _add_text(kp, keywords_case(art) or kw, 9)
    if kw and not kw.endswith("."):
        _add_text(kp, ".", 9)
    els.append(tbl._tbl)
    return els


def _add_link(par, url, text, color=None):
    from lxml import etree
    part = par.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                         is_external=True)
    hl = etree.SubElement(par._p, _qn("hyperlink"))
    hl.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", rid)
    if text:
        r = etree.SubElement(hl, _qn("r"))
        if color:
            rpr = etree.SubElement(r, _qn("rPr"))
            c = etree.SubElement(rpr, _qn("color"))
            c.set(_qn("val"), color)
            sz = etree.SubElement(rpr, _qn("sz"))
            sz.set(_qn("val"), "18")
        t = etree.SubElement(r, _qn("t"))
        t.text = text
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")


def style_document(doc):
    """Times New Roman, 9 pt iki yana yaslı gövde, dergi başlık biçimleri."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    def font(st, size, bold=None, italic=None):
        st.font.name = "Times New Roman"
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(_qn("rFonts"))
        if rfonts is None:
            from lxml import etree
            rfonts = etree.SubElement(rpr, _qn("rFonts"))
        for a in ("ascii", "hAnsi", "cs", "eastAsia"):
            rfonts.set(_qn(a), "Times New Roman")
        for a in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
            rfonts.attrib.pop(_qn(a), None)
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0, 0, 0)
        if bold is not None:
            st.font.bold = bold
        if italic is not None:
            st.font.italic = italic

    by_name = {st.name: st for st in doc.styles}
    names = set(by_name)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in names:
            st = by_name[name]
            font(st, 9)
            pf = st.paragraph_format
            pf.space_before, pf.space_after = Pt(0), Pt(3)
            pf.line_spacing = 1.0
            if name != "Normal":
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                pf.first_line_indent = Pt(10)
    for name, size, align in (("Heading 1", 10, WD_ALIGN_PARAGRAPH.CENTER), ("Heading 2", 10, None),
                              ("Heading 3", 10, None), ("Title", 12, WD_ALIGN_PARAGRAPH.CENTER)):
        if name in names:
            st = by_name[name]
            font(st, size, bold=True, italic=False)
            pf = st.paragraph_format
            pf.space_before, pf.space_after = Pt(8), Pt(4)
            pf.keep_with_next = True
            if align is not None:
                pf.alignment = align
    for name in ("Image Caption", "Table Caption", "Caption"):
        if name in names:
            st = by_name[name]
            font(st, 9, italic=False)
            st.paragraph_format.space_after = Pt(6)
            st.paragraph_format.alignment = (WD_ALIGN_PARAGRAPH.CENTER if name == "Image Caption"
                                             else WD_ALIGN_PARAGRAPH.LEFT)
    for name in ("Captioned Figure", "Figure"):
        if name in names:
            by_name[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ("Hyperlink",):
        if name in names:
            by_name[name].font.color.rgb = RGBColor(0, 0, 0)


def column_sections(doc):
    """İşaret paragraflarını sürekli bölüm sonlarına çevirir: geniş tablo/şekiller tek sütunda, metin iki sütunda.
    Bir bölüm sonu, kendinden önceki içeriğin sütun sayısını belirler."""
    from copy import deepcopy
    body = doc.element.body
    final = body.find(_qn("sectPr"))
    marks = {COLS_MARK[1]: 1, COLS_MARK[2]: 2}
    last_break = None   # en son bölüm sonu paragrafı
    for p in list(body.iterchildren(_qn("p"))):
        txt = "".join(t.text or "" for t in p.iter(_qn("t"))).strip()
        if txt not in marks:
            continue
        # TEKSUTUN işareti: buraya kadarki kısım tek sütun (ör. geniş tablo) bitti -> önceki içerik 1 sütun
        # IKISUTUN işareti: buraya kadarki iki sütunlu metin bitti -> önceki içerik 2 sütun
        cols = 1 if txt == COLS_MARK[1] else 2
        prev = p.getprevious()
        empty = prev is None or prev is last_break or (prev.tag == _qn("p") and prev.find(".//" + _qn("sectPr")) is not None)
        if empty:
            body.remove(p)
            continue
        for r in list(p):
            if r.tag != _qn("pPr"):
                p.remove(r)
        sp = deepcopy(final)
        for ref in sp.findall(_qn("headerReference")) + sp.findall(_qn("footerReference")) + sp.findall(_qn("titlePg")):
            sp.remove(ref)
        _set_cols(sp, cols, 360)
        t = sp.find(_qn("type"))
        if t is None:
            from lxml import etree
            t = etree.Element(_qn("type"))
            sp.insert(0, t)
        t.set(_qn("val"), "continuous")
        _ppr(p).append(sp)
        last_break = p
    return final


def _ppr(p):
    from lxml import etree
    ppr = p.find(_qn("pPr"))
    if ppr is None:
        ppr = etree.Element(_qn("pPr"))
        p.insert(0, ppr)
    return ppr


def page_layout(doc, art, opts, info, n_front):
    """Sayfa boyutu, ilk sayfa üst/alt bilgisi, diğer sayfalarda başlık satırı; ön kısım tek, gövde iki sütun."""
    from copy import deepcopy
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    from lxml import etree

    body = doc.element.body
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin, sec.bottom_margin = Inches(0.75), Inches(0.75)
    sec.header_distance, sec.footer_distance = Inches(0.3), Inches(0.3)
    sec.different_first_page_header_footer = True

    # ön kısmın sonuna tek sütunlu bölüm sonu
    final_sectPr = body.find(_qn("sectPr"))
    front_sect = deepcopy(final_sectPr)
    _set_cols(front_sect, 1)
    for ref in front_sect.findall(_qn("headerReference")) + front_sect.findall(_qn("footerReference")):
        front_sect.remove(ref)
    brk = doc.add_paragraph()._p
    ppr = brk.get_or_add_pPr()
    ppr.append(front_sect)
    body.remove(brk)
    children = [c for c in body if c.tag != _qn("sectPr")]
    children[n_front - 1].addnext(brk)
    _set_cols(final_sectPr, 2, 360)
    t = final_sectPr.find(_qn("type"))
    if t is None:
        t = etree.SubElement(final_sectPr, _qn("type"))
        final_sectPr.insert(0, t)
    t.set(_qn("val"), "continuous")
    doc.sections[-1].start_type = WD_SECTION.CONTINUOUS

    first = doc.sections[0]
    first.different_first_page_header_footer = True
    # ilk sayfa üst bilgisi: dergi adı + DOI (mavi kutu) ve logo
    fh = first.first_page_header
    ht = fh.add_table(rows=1, cols=2, width=Inches(7))
    c0, c1 = ht.cell(0, 0), ht.cell(0, 1)
    _grid(ht, (Inches(5.9), Inches(1.1)))
    _shade(c0, "C1D1EB")
    p = c0.paragraphs[0]
    _add_text(p, "Science Journal of University of Zakho", 9, False)
    p2 = c0.add_paragraph()
    _add_text(p2, "DOI: " + info["doi"], 9, italic=True)
    if info.get("logo") and os.path.isfile(info["logo"]):
        lp = c1.paragraphs[0]
        lp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lp.add_run().add_picture(info["logo"], width=Inches(1.0))
    if fh.paragraphs and not fh.paragraphs[0].text:
        fh.paragraphs[0]._p.getparent().remove(fh.paragraphs[0]._p)

    # ilk sayfa altlığı: dergi bilgi kutusu
    ff = first.first_page_footer
    ft = ff.add_table(rows=1, cols=3, width=Inches(7))
    _borders(ft, top=12, bottom=12)
    a, b, c = ft.cell(0, 0), ft.cell(0, 1), ft.cell(0, 2)
    _grid(ft, (Inches(1.5), Inches(2.6), Inches(2.9)))
    _add_text(a.paragraphs[0], "Access this article online", 8)
    if info.get("qr") and os.path.isfile(info["qr"]):
        a.add_paragraph().add_run().add_picture(info["qr"], width=Inches(0.9))
    bp = b.paragraphs[0]
    _add_text(bp, "DOI: ", 8)
    _add_link(bp, info["doi"], info["doi"], color="0000FF")
    _add_text(b.add_paragraph(), "Printed ISSN 2663-628X;", 8)
    _add_text(b.add_paragraph(), "Electronic ISSN 2663-6298", 8)
    _add_text(c.paragraphs[0], "Science Journal of University of Zakho", 8, bold=True)
    _add_text(c.add_paragraph(), "Vol. %s, No. %02d, pp. %s, %s-%s" % (opts.volume, int(opts.issue), info["pages"],
                                                                      opts.month, opts.year), 8, bold=True)
    _add_text(c.add_paragraph(), "This is an open access article under a CC BY-NC-SA 4.0 license "
                                 "(https://creativecommons.org/licenses/by-nc-sa/4.0/)", 8)
    pn = ff.paragraphs[0] if ff.paragraphs and not ff.paragraphs[0].text else ff.add_paragraph()
    pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(pn.add_run(), "PAGE")
    if ff.paragraphs[0]._p is not pn._p and not ff.paragraphs[0].text:
        ff.paragraphs[0]._p.getparent().remove(ff.paragraphs[0]._p)
    ff._element.append(pn._p)

    # diğer sayfalar
    for s in doc.sections:
        s.header.is_linked_to_previous = s is not first and True
        s.footer.is_linked_to_previous = s is not first and True
    hp = first.header.paragraphs[0]
    hp.text = ""
    _add_text(hp, info["runhead"], 10, italic=True)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(7), alignment=2)
    _add_text(hp, "\tSJUOZ|VOL%s|%s %s|P%s" % (opts.volume, opts.month[:3].upper(), opts.year, info["pages"]), 8,
              bold=True)
    fp = first.footer.paragraphs[0]
    fp.text = ""
    _add_text(fp, "DOI: ", 8, italic=True)
    _add_link(fp, info["doi"], info["doi"])
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(3.5), alignment=1)
    fp.add_run("\t")
    _field(fp.add_run(), "PAGE")
    if opts.first_page:
        pg = etree.SubElement(first._sectPr, _qn("pgNumType"))
        pg.set(_qn("start"), str(opts.first_page))
    return doc


def running_head(art):
    surnames = [au["name"].split()[-1] for au in art.authors if au["name"].split()]
    if len(surnames) >= 3:
        return surnames[0] + " et al."
    return " & ".join(surnames[:2])


def qr_from_pdf(pdf_path, png_path):
    """PDF'in ilk sayfasındaki QR kodunu resim olarak keser (PyMuPDF varsa)."""
    try:
        import pymupdf
    except ImportError:
        try:
            import fitz as pymupdf
        except ImportError:
            return None
    try:
        page = pymupdf.open(pdf_path)[0]
        hit = page.search_for("Access this article online")
        if not hit:
            return None
        r = hit[0]
        clip = pymupdf.Rect(r.x0 - 2, r.y1 + 1, r.x0 + 80, r.y1 + 80)
        page.get_pixmap(dpi=300, clip=clip).save(png_path)
        return png_path
    except Exception:  # noqa
        return None


def make_docx(art, writer, opts, outdir, docx_path, info):
    """Word çıktısını üretir; başarılıysa yol, değilse None."""
    import subprocess
    try:
        import docx
    except ImportError:
        print("Bilgi: Word çıktısı için 'pip install python-docx pypandoc_binary' gerekir; Word dosyası atlandı.")
        return None
    try:
        from .texinput import pandoc_exe
        exe = pandoc_exe()
    except SystemExit:
        print("Bilgi: Word çıktısı için 'pip install pypandoc_binary' gerekir; Word dosyası atlandı.")
        return None
    from .common import NO_WINDOW
    src = pandoc_latex("\n".join(writer.out))
    tmp_tex, tmp_docx = docx_path + ".tmp.tex", docx_path + ".tmp.docx"
    with open(tmp_tex, "w", encoding="utf-8") as f:
        f.write(src)
    try:
        res = subprocess.run([exe, tmp_tex, "-f", "latex", "-t", "docx", "-o", tmp_docx,
                              "--resource-path=%s" % outdir], capture_output=True, timeout=300, **NO_WINDOW)
        if res.returncode != 0 or not os.path.isfile(tmp_docx):
            warn("Word çıktısı üretilemedi (pandoc): %s" % res.stderr.decode("utf-8", "replace")[:200])
            return None
        d = docx.Document(tmp_docx)
    finally:
        for f in (tmp_tex, tmp_docx):
            if os.path.isfile(f):
                os.remove(f)
    style_document(d)
    els = build_front(d, art, opts, info)
    anchor = d.element.body[0]
    for el in els:
        anchor.addprevious(el)
    page_layout(d, art, opts, info, len(els))
    column_sections(d)
    d.core_properties.title = art.title
    d.core_properties.author = "; ".join(a["name"] for a in art.authors)
    d.save(docx_path)
    return docx_path
